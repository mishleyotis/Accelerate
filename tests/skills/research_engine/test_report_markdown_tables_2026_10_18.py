"""A markdown table in a section body reaches the Doc as a Word table, and the
gold gate refuses a Doc that carries one as raw pipe text (EPIC, 2026-10-18:
the delivered assessment report held 74 literal `| :-: |` separator lines and
one real table, and GS-RPT printed PASS)."""
from docx import Document

from engine import gold_standard as GS, reports

BODY = """## 4.1 Pillar results

| Pillar | Score | Cohort average |
| :-: | :-: | :-: |
| P1 | 1.44 | 1.38 |
| P3 | 1.06 | 1.39 |

P3 trails by 0.33."""


def test_a_markdown_table_in_a_body_renders_as_a_word_table():
    doc = Document()
    reports._emit_body(doc, BODY)
    assert len(doc.tables) == 1
    t = doc.tables[0]
    assert [c.text for c in t.rows[0].cells] == ["Pillar", "Score", "Cohort average"]
    assert [c.text for c in t.rows[2].cells] == ["P3", "1.06", "1.39"]
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "| :-: |" not in text and "P3 trails by 0.33." in text


def test_prose_with_a_pipe_is_not_mistaken_for_a_table():
    doc = Document()
    reports._emit_body(doc, "Score | weight is not a table.\nSecond line.")
    assert not doc.tables
    assert "Score | weight" in doc.paragraphs[0].text


def test_a_ragged_row_is_padded_not_dropped():
    doc = Document()
    reports._emit_lines(doc, ["| A | B | C |", "| --- | --- | --- |", "| 1 | 2 |"])
    assert [c.text for c in doc.tables[0].rows[1].cells] == ["1", "2", ""]


def test_the_gold_gate_refuses_tables_rendered_as_raw_text(tmp_path):
    raw = tmp_path / "raw.docx"
    d = Document()
    for line in BODY.splitlines():
        d.add_paragraph(line)
    d.save(str(raw))
    codes = {f["code"] for f in GS.report_findings(raw, kind="assessment")}
    assert "GS-RPT-TABLES" in codes

    good = tmp_path / "good.docx"
    d2 = Document()
    reports._emit_body(d2, BODY)
    d2.save(str(good))
    assert "GS-RPT-TABLES" not in {f["code"] for f in GS.report_findings(good, kind="assessment")}
