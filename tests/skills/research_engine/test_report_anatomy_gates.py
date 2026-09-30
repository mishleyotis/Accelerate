"""The gold gate reads the report's ANATOMY, not only its volume.

Owner, 2026-09-07: "huge discrepancy with the formatting … even the cover
page is off … a lot of placeholder and unnecessary text … Did you check how
the golden standard was written?" It had not been: the gate of the day was
calibrated to the reference's table COUNT, not to how its sections are
composed. Both golden reports were measured (gold_reference.json
`section_tables`, `cover_labels`, `front_matter_h1`) and these gates hold a
report to that anatomy:

    GS-RPT-COVER                 a boxed title + a metadata grid carrying the
                                 Doc's identity labels; a bare Title is not a cover
    GS-RPT-FRONTMATTER           Contents, then Document Control and Catalogue Binding
    GS-RPT-SECTION-DISTRIBUTION  no numbered section barren where the reference tabulates
    GS-RPT-DEGENERATE-TABLE      a table whose only varying column is the label
    GS-RPT-PROSE-DUMP            a field register typed as one `;`-separated sentence

Every floor is at or below what the reference itself meets, the renderer's
own output clears all five, and a good fixture mutated one way at a time
proves each gate fires.
"""
from __future__ import annotations

import json
import zipfile

import pytest
from docx import Document

from engine import gold_standard as GS
from engine import report_spec as RS
from engine import reports as R
from engine import rubric
from engine import template as T

from fixtures import bank_evidence, scored_run, sign_off_sections, write_report

_V2 = ("GS-RPT-COVER", "GS-RPT-FRONTMATTER", "GS-RPT-SECTION-DISTRIBUTION",
       "GS-RPT-DEGENERATE-TABLE", "GS-RPT-PROSE-DUMP")


@pytest.fixture(scope="module")
def gold() -> dict:
    return json.loads((T.TEMPLATES_DIR / "gold_reference.json").read_text())


def _codes(path, kind="assessment", subcaps=690):
    return {f["code"] for f in GS.report_findings(path, kind=kind, subcaps=subcaps)}


# ── calibration: no anatomy floor exceeds what the reference itself meets ──

def test_section_floors_never_exceed_the_reference_section_counts(gold):
    """The same discipline as every other floor: a per-section floor the
    reference would fail is a floor nobody measured."""
    for kind in ("research", "assessment"):
        anat = GS.section_floors(kind)
        ref = gold["reports"][kind]["section_tables"]
        assert anat["section_floors"], kind
        for num, floor in anat["section_floors"].items():
            assert 1 <= floor <= int(ref[num]), (kind, num, floor, ref[num])


def test_the_section_floors_cover_every_pinned_section_and_scale_down(gold):
    for kind, key in (("research", "client_research"), ("assessment", "assessment")):
        ref = gold["reports"][kind]["section_tables"]
        assert set(ref) == {s.id for s in RS.SPECS[key].sections}, kind
        small = GS.section_floors(kind, subcaps=6)["section_floors"]
        full = GS.section_floors(kind)["section_floors"]
        assert all(small[n] <= full[n] for n in ref)
        assert all(v >= 1 for v in small.values())


def test_the_reference_front_matter_and_cover_labels_are_recorded(gold):
    for kind in ("research", "assessment"):
        anat = GS.section_floors(kind)
        h1 = gold["reports"][kind]["heading1"]
        for want in anat["front_matter_h1"]:
            assert want in h1, (kind, want)
        assert "Surface Alignment" not in h1
        assert anat["cover_labels"], kind


# ── fixture: a full report with cover, front matter and per-section tables ─

def _full_report(path, *, kind="assessment", degenerate=False, dump=False,
                 barren_section=None, cover=True, front=True):
    """A report of the reference's SHAPE — cover box + metadata grid +
    Contents + Document Control, then each numbered section carrying its
    reference table count — so an anatomy finding can be provoked or
    avoided on purpose."""
    import docx
    d = docx.Document()
    labels = GS.section_floors(kind)["cover_labels"]
    if cover:
        box = d.add_table(rows=1, cols=1)
        box.rows[0].cells[0].text = "Acme Credit Union"
        grid = d.add_table(rows=len(labels), cols=2)
        for i, lb in enumerate(labels):
            grid.rows[i].cells[0].text = lb
            grid.rows[i].cells[1].text = f"{lb.lower()} value {i}"
    if front:
        d.add_paragraph("Contents", style="Heading 1")
        d.add_paragraph("Document Control and Catalogue Binding", style="Heading 1")
        dc = d.add_table(rows=2, cols=3)
        for i, c in enumerate(("Field", "Value", "Resolution source")):
            dc.rows[0].cells[i].text = c
        dc.rows[1].cells[0].text = "Catalogue version"
        dc.rows[1].cells[1].text = "v7.0"
        dc.rows[1].cells[2].text = "Catalogue_Meta!version"
    section_tables = GS.gold_reference()["reports"][kind]["section_tables"]
    d.add_paragraph(" ".join(f"[E-{i}]" for i in range(1, 200)))
    d.add_paragraph(" ".join(["word"] * 12000) + " coverage unknown")
    for num, count in sorted(section_tables.items(), key=lambda kv: int(kv[0])):
        d.add_paragraph(f"{num}. Section {num}", style="Heading 1")
        n = 0 if (barren_section == num) else count
        for _ in range(n):
            t = d.add_table(rows=1, cols=3)
            for i, c in enumerate(("Cell", "Score", "Evidence")):
                t.rows[0].cells[i].text = c
            for r in range(3):
                cells = t.add_row().cells
                cells[0].text = f"P{r}C1"
                cells[1].text = f"{r + 1}.0"
                cells[2].text = f"E-{r}"
    if degenerate:
        t = d.add_table(rows=1, cols=3)
        for i, c in enumerate(("Field", "State", "Route")):
            t.rows[0].cells[i].text = c
        for f in ("website", "employees", "assets", "branches"):
            cells = t.add_row().cells
            cells[0].text = f
            cells[1].text = "STATED"     # constant
            cells[2].text = ""            # empty
    if dump:
        d.add_paragraph(
            "All ten fields are STATED: website acme.example ([E-1], High); "
            "employees 265 ([E-1], Medium); assets $1.18B ([E-2], High); "
            "branches 16 ([E-1], High); founded 1955 ([E-1], High); regulator "
            "NCUA ([E-2], High).")
    d.save(str(path))
    with zipfile.ZipFile(str(path), "a") as z:
        z.writestr("word/header1.xml", "<hdr/>")
    return path


def test_a_full_reference_shaped_report_raises_no_anatomy_finding(tmp_path):
    for kind in ("research", "assessment"):
        p = _full_report(tmp_path / f"{kind}.docx", kind=kind)
        codes = _codes(p, kind=kind)
        assert not (codes & set(_V2)), (kind, codes & set(_V2))


def test_a_bare_title_with_no_cover_grid_is_refused(tmp_path):
    p = _full_report(tmp_path / "nocover.docx", cover=False)
    assert "GS-RPT-COVER" in _codes(p)


def test_a_cover_missing_a_required_label_is_refused(tmp_path):
    import docx
    d = docx.Document()
    box = d.add_table(rows=1, cols=1)
    box.rows[0].cells[0].text = "Acme Credit Union"
    grid = d.add_table(rows=1, cols=2)
    grid.rows[0].cells[0].text = "ASSESSMENT ID"     # one label, the rest missing
    grid.rows[0].cells[1].text = "DMA-X"
    d.add_paragraph("Contents", style="Heading 1")
    d.add_paragraph("1. Section 1", style="Heading 1")
    d.save(str(tmp_path / "partial.docx"))
    with zipfile.ZipFile(str(tmp_path / "partial.docx"), "a") as z:
        z.writestr("word/header1.xml", "<hdr/>")
    assert "GS-RPT-COVER" in _codes(tmp_path / "partial.docx")


def test_missing_document_control_front_matter_is_refused(tmp_path):
    p = _full_report(tmp_path / "nofront.docx", front=False)
    assert "GS-RPT-FRONTMATTER" in _codes(p)


def test_a_barren_section_is_refused(tmp_path):
    """Section 8 of the reference carries 54 tables; a run that leaves it
    empty while the total still clears is hiding a prose section behind a
    rich one."""
    p = _full_report(tmp_path / "barren.docx", barren_section="8")
    assert "GS-RPT-SECTION-DISTRIBUTION" in _codes(p)


def test_a_degenerate_table_is_refused(tmp_path):
    p = _full_report(tmp_path / "degen.docx", degenerate=True)
    assert "GS-RPT-DEGENERATE-TABLE" in _codes(p)


def test_a_table_with_one_varying_column_is_not_degenerate():
    # the reference's identity-check table: Result constant PASS, Basis varies.
    ok = [["Check", "Result", "Basis"],
          ["name matches", "PASS", "E-1, E-2"],
          ["regulator", "PASS", "E-1, E-5"],
          ["footprint", "PASS", "E-3"],
          ["charter", "PASS", "E-2"]]
    assert not GS._degenerate_table(ok)
    bad = [["Field", "State", "Route"],
           ["website", "STATED", ""],
           ["employees", "STATED", ""],
           ["assets", "STATED", ""],
           ["branches", "STATED", ""]]
    assert GS._degenerate_table(bad)
    # a two-row table is too small to judge — the cover grid is one
    assert not GS._degenerate_table(bad[:3])


def test_a_field_register_typed_as_a_paragraph_is_refused(tmp_path):
    p = _full_report(tmp_path / "dump.docx", dump=True)
    assert "GS-RPT-PROSE-DUMP" in _codes(p)


def test_interpretive_prose_with_inline_citations_is_not_a_dump():
    """The reference cites inline in whole sentences; only a short cited
    field entry counts, and only many of them in one paragraph is a dump."""
    argued = ("Golden 1 has assembled a modern rails-and-platform foundation "
              "and, in its own 2024 discovery [E-005], named the value it wants "
              "to convert next. The core runs on Fiserv DNA; real-time payments "
              "went live on RTP and FedNow in fall 2024 [E-021]; a custom Zest "
              "AI scorecard lifted protected-class approvals by 28% [E-055].")
    assert GS._prose_dump_clauses(argued) < GS.PROSE_DUMP_CLAUSES


# ── the renderer's own cover and front matter clear the gates ──────────────

class _Stub:
    def __init__(self, rows):
        self._rows = rows

    def rows(self, name):
        return self._rows if name == "Pillar_Summary" else []


def test_the_cover_level_is_the_sheets_own_label_or_the_rubrics_never_typed():
    """The OVERALL row's `Maturity` is what the package gate reconciles to;
    when the sheet carries the score without its label the rubric names it
    (`rubric.maturity_level`) — the four display bands and the 1–5 score
    scale are the two axes the charter keeps apart, and no level word is
    typed in the renderer."""
    labelled = _Stub([{"Pillar": "OVERALL", "Score": 2.25, "Maturity": "M2"}])
    assert R._overall_maturity(labelled) == ("2.25", "M2")
    for score in (1.0, 1.5, 2.49, 3.5, 4.5, 5.0):
        bare = _Stub([{"Pillar": "OVERALL", "Score": score, "Maturity": ""}])
        assert R._overall_maturity(bare) == (str(score), rubric.maturity_level(score))
    assert R._overall_maturity(_Stub([{"Pillar": "OVERALL", "Score": None}])) == ("", "")
    assert R._overall_maturity(_Stub([])) == ("", "")


def test_the_rendered_reports_open_with_the_docs_cover_and_front_matter(tmp_path):
    run, wb, cells, ev = scored_run(tmp_path)
    eids = bank_evidence(wb, cells[0], n=7)
    for key in RS.SPECS:
        write_report(wb, key, eids, run=run)
    sign_off_sections(wb)
    for key, kind in (("assessment", "assessment"), ("client_research", "research")):
        out = R.render(wb, RS.SPECS[key], tmp_path / "out")
        layout = GS._docx_layout(out["path"])
        anat = GS.section_floors(kind, subcaps=len(cells))
        # a boxed title, the metadata grid and the binding table sit before §1
        assert len(layout["cover_tables"]) == R.FRONT_MATTER_TABLES, key
        flat = GS._flat([r for t in layout["cover_tables"] for r in t])
        for lb in anat["cover_labels"]:
            assert lb.casefold() in flat, (key, lb)
        assert "acme credit union" in flat
        assert layout["front_h1"] == anat["front_matter_h1"], layout["front_h1"]
        doc = Document(out["path"])
        h1 = [p.text for p in doc.paragraphs if p.style.name in ("Heading 1", "Title")]
        assert "Surface Alignment" not in h1
        assert not any(p.style.name == "Title" for p in doc.paragraphs)
        # every numbered section tabulates where the reference does
        for num, floor in anat["section_floors"].items():
            assert layout["sections"].get(num, 0) >= floor, (key, num, layout["sections"])
        codes = _codes(out["path"], kind=kind, subcaps=len(cells))
        assert not (codes & set(_V2)), (key, codes & set(_V2))
        if key == "assessment":
            # the cover states the OVERALL row's own score and label
            score, level = R._overall_maturity(wb)
            assert score and level
            assert f"{score} of 5.0 ({level})".casefold() in flat
