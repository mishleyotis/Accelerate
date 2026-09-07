"""A report can pass every volume floor and still be the wrong SHAPE.

THE DEFECT THESE TESTS PIN, measured 2026-09-06 on a delivered pair of reports
that the gate of the day passed with zero findings:

    assessment   50 tables against the reference's 92, paragraph words 1.57x
    research     26 tables against the reference's 39, paragraph words 1.57x

The owner's report was "the 2 reports lack depth ... do not adhere to template
requirements eg where tables are, I see paragraphs." The gate counted words,
citations and section headings, none of which can tell a report that TABULATES
its register from one that describes it in a paragraph — and because more prose
raises a word count, the defect was not merely uncaught, it was REWARDED.

The root cause was in the renderer, not the writer: `reports._emit_body` turned
every body line into a paragraph, so the only tables a report could contain were
the whole sheets `_tables_for` emits per declared input (~44 for the assessment).
Golden 1 carries 92, so roughly half of the reference's tables are AUTHORED
inside section bodies — which the renderer had no way to produce. A writer asked
for a table could not make one.

So three things are asserted here: the new gates catch the shape, the renderer
can now render an authored table, and the anti-patterns reach the writer before
it writes (engine/authoring.py) rather than only the gate after it ships.
"""
from __future__ import annotations

import json
import zipfile

import pytest

from engine import authoring as A
from engine import gold_standard as GS
from engine import report_spec as RS
from engine import reports as R
from engine import template as T


@pytest.fixture(scope="module")
def gold() -> dict:
    return json.loads((T.TEMPLATES_DIR / "gold_reference.json").read_text())


# ── the calibration discipline, extended to the structure floors ──────────

def test_the_table_floor_is_at_or_below_what_golden_1_meets(gold):
    """The rule gold_standard.py has stated since it was written: a floor the
    reference itself would fail is a floor nobody measured."""
    for kind in ("research", "assessment"):
        floors = GS.depth_floors(kind)
        assert floors["tables"] <= gold["reports"][kind]["tables"], (kind, floors)


def test_the_prose_and_dump_limits_do_not_refuse_the_reference(gold):
    for kind in ("research", "assessment"):
        ref = gold["reports"][kind]
        floors = GS.depth_floors(kind)
        # The reference's own prose sits inside the inflation limit …
        assert ref["words_paragraphs"] <= (
            floors["reference_paragraph_words"] * GS.PROSE_INFLATION_LIMIT)
        # … and its own average table is, by construction, 1x its average.
        assert GS._gold_avg_table_words(kind) >= 1


def test_the_table_floor_scales_with_the_run_and_never_vanishes():
    small = GS.depth_floors("assessment", subcaps=6)
    full = GS.depth_floors("assessment")
    assert 1 <= small["tables"] < full["tables"]


# ── fixtures: a report of a chosen shape ──────────────────────────────────

def _report(path, *, tables, rows_per_table, prose_words, kind="assessment"):
    """A synthetic report carrying an exact number of tables of an exact size,
    and an exact quantity of prose — so a shape finding can be provoked or
    avoided on purpose."""
    import docx
    d = docx.Document()
    for h in RS.numbered_headings(kind if kind == "assessment" else "client_research"):
        d.add_paragraph(h, style="Heading 1")
    # citations, so the volume gates are not what fires
    d.add_paragraph(" ".join(f"[E-{i}]" for i in range(1, 200)))
    d.add_paragraph(" ".join(["word"] * prose_words))
    for _ in range(tables):
        t = d.add_table(rows=1, cols=3)
        for i, c in enumerate(("A", "B", "C")):
            t.rows[0].cells[i].text = c
        for _ in range(rows_per_table):
            cells = t.add_row().cells
            for i in range(3):
                cells[i].text = "cell value here"
    d.save(str(path))
    with zipfile.ZipFile(str(path), "a") as z:
        z.writestr("word/header1.xml", "<hdr/>")
    return path


def _codes(path, kind="assessment", subcaps=690):
    return {f["code"] for f in GS.report_findings(path, kind=kind, subcaps=subcaps)}


# ── detection ─────────────────────────────────────────────────────────────

def test_prose_where_a_table_belongs_is_refused(tmp_path):
    """Few tables, inflated prose — the delivered pair's exact shape."""
    p = _report(tmp_path / "thin.docx", tables=50, rows_per_table=2,
                prose_words=20000)
    codes = _codes(p)
    assert "GS-RPT-TABLES" in codes
    assert "GS-RPT-PROSE-FOR-STRUCTURE" in codes


def test_a_sheet_emitted_whole_is_refused(tmp_path):
    """Enough tables, but each one is a pasted sheet rather than an extract."""
    p = _report(tmp_path / "dump.docx", tables=92, rows_per_table=400,
                prose_words=11000)
    assert "GS-RPT-TABLE-DUMP" in _codes(p)


def test_a_report_of_the_reference_shape_passes_every_structure_gate(tmp_path):
    """The floors must be reachable: a report built to Golden 1's own measured
    shape raises none of the three structure findings."""
    p = _report(tmp_path / "good.docx", tables=92, rows_per_table=3,
                prose_words=11000)
    codes = _codes(p)
    for code in ("GS-RPT-TABLES", "GS-RPT-PROSE-FOR-STRUCTURE",
                 "GS-RPT-TABLE-DUMP"):
        assert code not in codes, (code, codes)


def test_long_prose_alone_is_not_refused_when_the_tables_are_there(tmp_path):
    """The finding is prose STANDING IN FOR structure, not prose. A section
    that tabulates what it should and still argues at length is not the
    defect, and a gate that punished it would push producers to cut analysis."""
    p = _report(tmp_path / "rich.docx", tables=92, rows_per_table=3,
                prose_words=30000)
    assert "GS-RPT-PROSE-FOR-STRUCTURE" not in _codes(p)


# ── the root cause: a writer can now author a table ───────────────────────

def test_an_authored_pipe_table_renders_as_a_real_table():
    """Before 2026-09-06 this rendered as a paragraph of pipe characters, so a
    writer asked for a table could not produce one and wrote prose instead."""
    import docx
    d = docx.Document()
    R._emit_body(d, "## Capability scorecard\n\n"
                    "| Cell | Score | Peer |\n"
                    "|---|---|---|\n"
                    "| P4C1.2.1 | 1.6 | 3.0 |\n"
                    "| P4C1.2.2 | 1.25 | 3.0 |\n\n"
                    "The scorecard above states the two cells this rests on.")
    assert len(d.tables) == 1
    t = d.tables[0]
    assert [c.text for c in t.rows[0].cells] == ["Cell", "Score", "Peer"]
    assert t.rows[1].cells[0].text == "P4C1.2.1"
    # the prose after the table is still a paragraph, not swallowed by it
    assert any("states the two cells" in p.text for p in d.paragraphs)


def test_prose_that_merely_contains_a_pipe_is_still_a_paragraph():
    import docx
    d = docx.Document()
    R._emit_body(d, "The estate splits OPS | CUST | DATA across four layers.")
    assert len(d.tables) == 0
    assert any("four layers" in p.text for p in d.paragraphs)


# ── the anti-patterns reach the writer BEFORE it writes ───────────────────

def test_every_antipattern_names_a_gate_and_a_repair():
    pats = A.antipatterns()
    assert pats, "the anti-pattern register must not be empty"
    for a in pats:
        for field in ("id", "gate", "what", "measured", "repair"):
            assert a.get(field), (a.get("id"), field)


def test_the_register_covers_every_structure_gate():
    """A gate with no entry in the register is a rule the writer is never told
    about until it refuses them."""
    gates = {a["gate"] for a in A.antipatterns()}
    for code in ("GS-RPT-TABLES", "GS-RPT-PROSE-FOR-STRUCTURE",
                 "GS-RPT-TABLE-DUMP", "GS-RPT-COVER", "GS-RPT-FRONTMATTER",
                 "GS-RPT-SECTION-DISTRIBUTION", "GS-RPT-DEGENERATE-TABLE",
                 "GS-RPT-PROSE-DUMP"):
        assert code in gates, (code, gates)


def test_the_brief_binds_the_antipatterns_to_this_sections_own_tables():
    """A general rule is advice; the section's own declared inputs are an
    instruction. §6 of the assessment declares the technology estate sheets."""
    text = A.brief("assessment", "6")
    assert "Benchmark and Technology Estate" in text
    assert "TABLES THIS SECTION OWES" in text
    for sheet in A.table_inputs(
            next(s for s in RS.SPECS["assessment"].sections if str(s.id) == "6")):
        assert sheet in text
    for a in A.antipatterns():
        assert a["id"] in text


def test_the_brief_names_an_empty_declared_input_as_the_thing_to_fix_first():
    """An input that carries no rows renders no table, and the section then
    'lacks a table' for a reason no amount of writing can repair."""
    class _WB:
        def rows(self, name):
            return [] if name == "Tech_Register" else [{"x": 1}]

    text = A.brief("assessment", "6", _WB())
    assert "Tech_Register" in text
    assert "EMPTY" in text
    ob = A.section_obligations("assessment", "6", _WB())
    assert "Tech_Register" in ob["tables_empty"]
    assert "Tech_Register" not in ob["tables_that_will_render"]


def test_preflight_reports_the_whole_reports_table_position():
    class _WB:
        def rows(self, name):
            return [{"x": 1}]

    out = A.preflight("assessment", _WB())
    assert out["tables_declared"] > 0
    assert out["tables_that_will_render"] == out["tables_declared"]
    assert out["empty_declared_inputs"] == []
    assert len(out["sections"]) == len(RS.SPECS["assessment"].sections)


# ══════════════════════════════════════════════════════════════════════════
# SAFEGUARDS v2 — cover, front matter, distribution, degenerate table,
# prose dump. Owner 2026-09-07: "huge discrepancy with the formatting ... even
# the cover page is off ... a lot of placeholder and unnecessary text ... Did
# you check how the golden standard was written?" The volume and total-table
# gates above pass a report whose cover is a bare heading, whose tables pile
# into one section, or which types a field register as a paragraph. These
# gates read the reference's own per-section anatomy (measured into
# gold_reference.json) and hold the report to it.
# ══════════════════════════════════════════════════════════════════════════

_V2 = ("GS-RPT-COVER", "GS-RPT-FRONTMATTER", "GS-RPT-SECTION-DISTRIBUTION",
       "GS-RPT-DEGENERATE-TABLE", "GS-RPT-PROSE-DUMP")


# ── calibration: no v2 floor exceeds what the reference itself meets ───────

def test_section_floors_never_exceed_the_reference_section_counts(gold):
    """The same discipline as the table floor: a per-section floor the
    reference would fail is a floor nobody measured."""
    for kind in ("research", "assessment"):
        anat = GS.section_floors(kind)
        ref = gold["reports"][kind]["section_tables"]
        for num, floor in anat["section_floors"].items():
            assert floor <= int(ref[num]), (kind, num, floor, ref[num])
            assert floor >= 1


def test_the_reference_front_matter_and_cover_labels_are_recorded(gold):
    for kind in ("research", "assessment"):
        anat = GS.section_floors(kind)
        # the two unnumbered H1s the reference opens with are a subset of its
        # measured heading1 list
        h1 = gold["reports"][kind]["heading1"]
        for want in anat["front_matter_h1"]:
            assert want in h1, (kind, want)
        assert anat["cover_labels"], kind


# ── fixture: a full report with cover, front matter and per-section tables ─

def _full_report(path, *, kind="assessment", degenerate=False, dump=False,
                 barren_section=None, cover=True, front=True):
    """A report of the reference's SHAPE — cover box + metadata grid + Contents
    + Document Control, then each numbered section carrying its reference table
    count — so a v2 finding can be provoked or avoided on purpose."""
    import docx
    d = docx.Document()
    labels = GS.section_floors(kind)["cover_labels"]
    if cover:
        box = d.add_table(rows=1, cols=1)
        box.rows[0].cells[0].text = "REV Federal Credit Union"
        grid = d.add_table(rows=len(labels), cols=2)
        for i, lb in enumerate(labels):
            grid.rows[i].cells[0].text = lb
            grid.rows[i].cells[1].text = f"{lb.lower()} value {i}"
    if front:
        d.add_paragraph("Contents", style="Heading 1")
        d.add_paragraph("Document Control and Catalogue Binding", style="Heading 1")
        dc = d.add_table(rows=2, cols=3)
        for i, c in enumerate(("Field", "Value", "Resolution source")):
            dc.rows[0].cells[i].text = c
        dc.rows[1].cells[0].text = "Catalogue version"
        dc.rows[1].cells[1].text = "v7.0"
        dc.rows[1].cells[2].text = "Catalogue_Meta!version"
    ref = json.loads((T.TEMPLATES_DIR / "gold_reference.json").read_text())
    section_tables = ref["reports"][kind]["section_tables"]
    d.add_paragraph(" ".join(f"[E-{i}]" for i in range(1, 200)))
    d.add_paragraph(" ".join(["word"] * 12000) + " coverage unknown")
    for num, count in sorted(section_tables.items(), key=lambda kv: int(kv[0])):
        d.add_paragraph(f"{num}. Section {num}", style="Heading 1")
        n = 0 if (barren_section == num) else count
        for _ in range(n):
            t = d.add_table(rows=1, cols=3)
            for i, c in enumerate(("Cell", "Score", "Evidence")):
                t.rows[0].cells[i].text = c
            for r in range(3):
                cells = t.add_row().cells
                cells[0].text = f"P{r}C1"
                cells[1].text = f"{r + 1}.0"
                cells[2].text = f"E-{r}"
    if degenerate:
        t = d.add_table(rows=1, cols=3)
        for i, c in enumerate(("Field", "State", "Route")):
            t.rows[0].cells[i].text = c
        for f in ("website", "employees", "assets", "branches"):
            cells = t.add_row().cells
            cells[0].text = f
            cells[1].text = "STATED"     # constant
            cells[2].text = ""            # empty
    if dump:
        d.add_paragraph(
            "All ten fields are STATED: website revfcu.com ([E-1], High); "
            "employees 265 ([E-1], Medium); assets $1.18B ([E-2], High); "
            "branches 16 ([E-1], High); founded 1955 ([E-1], High); regulator "
            "NCUA ([E-2], High).")
    d.save(str(path))
    with zipfile.ZipFile(str(path), "a") as z:
        z.writestr("word/header1.xml", "<hdr/>")
    return path


def test_a_full_reference_shaped_report_raises_no_v2_finding(tmp_path):
    for kind in ("research", "assessment"):
        p = _full_report(tmp_path / f"{kind}.docx", kind=kind)
        codes = _codes(p, kind=kind)
        assert not (codes & set(_V2)), (kind, codes & set(_V2))


def test_a_bare_title_with_no_cover_grid_is_refused(tmp_path):
    p = _full_report(tmp_path / "nocover.docx", cover=False)
    assert "GS-RPT-COVER" in _codes(p)


def test_a_cover_missing_a_required_label_is_refused(tmp_path):
    import docx
    d = docx.Document()
    box = d.add_table(rows=1, cols=1)
    box.rows[0].cells[0].text = "REV Federal Credit Union"
    grid = d.add_table(rows=1, cols=2)
    grid.rows[0].cells[0].text = "ASSESSMENT ID"     # only one label, missing the rest
    grid.rows[0].cells[1].text = "DMA-X"
    d.add_paragraph("Contents", style="Heading 1")
    d.add_paragraph("1. Section 1", style="Heading 1")
    d.save(str(tmp_path / "partial.docx"))
    with zipfile.ZipFile(str(tmp_path / "partial.docx"), "a") as z:
        z.writestr("word/header1.xml", "<hdr/>")
    assert "GS-RPT-COVER" in _codes(tmp_path / "partial.docx")


def test_missing_document_control_front_matter_is_refused(tmp_path):
    p = _full_report(tmp_path / "nofront.docx", front=False)
    assert "GS-RPT-FRONTMATTER" in _codes(p)


def test_a_barren_section_is_refused(tmp_path):
    """Section 8 of the reference carries 54 tables; a run that leaves it empty
    while the total still clears is hiding a prose section behind a rich one."""
    p = _full_report(tmp_path / "barren.docx", barren_section="8")
    assert "GS-RPT-SECTION-DISTRIBUTION" in _codes(p)


def test_a_degenerate_table_is_refused(tmp_path):
    p = _full_report(tmp_path / "degen.docx", degenerate=True)
    assert "GS-RPT-DEGENERATE-TABLE" in _codes(p)


def test_a_table_with_one_varying_column_is_not_degenerate():
    # the reference's identity-check table: Result constant PASS, Basis varies.
    ok = [["Check", "Result", "Basis"],
          ["name matches", "PASS", "E-1, E-2"],
          ["regulator", "PASS", "E-1, E-5"],
          ["footprint", "PASS", "E-3"],
          ["charter", "PASS", "E-2"]]
    assert not GS._degenerate_table(ok)
    bad = [["Field", "State", "Route"],
           ["website", "STATED", ""],
           ["employees", "STATED", ""],
           ["assets", "STATED", ""],
           ["branches", "STATED", ""]]
    assert GS._degenerate_table(bad)


def test_a_field_register_typed_as_a_paragraph_is_refused(tmp_path):
    p = _full_report(tmp_path / "dump.docx", dump=True)
    assert "GS-RPT-PROSE-DUMP" in _codes(p)


def test_interpretive_prose_with_inline_citations_is_not_a_dump():
    """The reference cites inline in whole sentences; only a short cited field
    entry counts, and only many of them in one paragraph is a dump."""
    argued = ("Golden 1 has assembled a modern rails-and-platform foundation "
              "and, in its own 2024 discovery [E-005], named the value it wants "
              "to convert next. The core runs on Fiserv DNA; real-time payments "
              "went live on RTP and FedNow in fall 2024 [E-021]; a custom Zest "
              "AI scorecard lifted protected-class approvals by 28% [E-055].")
    assert GS._prose_dump_clauses(argued) < GS.PROSE_DUMP_CLAUSES
