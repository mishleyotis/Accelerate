"""The report gates read the section, not just its length and its ids.

MEASURED 2026-09-08 on an engine-built package: a body consisting of ONE
paragraph pasted 73 times across both reports — every block heading present,
every countable id present — cleared `narrative.write`, `narrative.review`,
`reports.render`, `assemble verify` and `engine.gold_standard package` with
zero findings. A REC card with all its prose under `## Root cause` and seven
bare block headings was accepted. A markdown table in a body landed as one
paragraph of pipe characters, while the renderer appended the 690-row
`Subcap_Scores` sheet six times and `Peer_Benchmarks` seven times, and every
one of those table words counted toward the 8,400-word LENGTH floor. The
brief handed the producers `--body-file` (the CLI takes `--json`) and the
validator `--verdict READY` (the CLI takes PASS|REVISE|FAIL). Golden 1, the
reference the gate is calibrated to, has none of these shapes: 92 curated
Doc tables of at most 22 rows, each inside the block that argues it, and no
paragraph repeated.

Every test here fails on the code as it stood that morning.
"""
from __future__ import annotations

import re
from pathlib import Path

import pytest
from docx import Document

from engine import brief, gold_standard as GS, grains, narrative as N
from engine import report_spec as RS, reports as R

from fixtures import (bank_evidence, report_ready_run, scored_run,
                      section_record, sign_off_sections, write_report)

OLD_FILLER = (
    "The public record for this institution is read here against the "
    "question the block asks, and the reading is stated so a reader can "
    "disagree with it rather than accept it. Nothing in this paragraph rests "
    "on a source that is not in the run's own register, and every figure it "
    "carries can be reopened from the excerpt that supplied it rather than "
    "recalled from anywhere else in the record of the engagement. ")


@pytest.fixture(scope="module")
def ready(tmp_path_factory):
    """One scored, shippable run for the module — building one takes
    minutes, and every test here only WRITES to Report_Narrative."""
    run = report_ready_run(tmp_path_factory.mktemp("depth"))
    return run


def _eids(run, n=8):
    wb = run.open()
    return list(wb.evidence_index())[:n]


# ── the writer reads the body ──────────────────────────────────────────────

def test_the_old_pasted_paragraph_body_is_refused(ready):
    wb = ready.open()
    eids = _eids(ready)
    sec = RS.SPECS["client_research"].section("3")
    body = "\n".join(f"## {b}\n" + OLD_FILLER * 3 + "\nSources: "
                     + " ".join(f"[{e}]" for e in eids) + "\n"
                     for b in sec.blocks)
    rec = section_record("3", eids, "client_research", Body=body)
    with pytest.raises(N.NarrativeRefusal, match="form-filling|more than once"):
        N.write(wb, "client_research", "3", rec, actor="probe", run=ready)


def test_a_block_heading_with_nothing_under_it_is_refused(ready):
    wb = ready.open()
    eids = _eids(ready)
    rec = section_record("8", eids, "assessment")
    lines = rec["Body"].split("\n")
    heads = [l for l in lines if l.startswith("## ")]
    rest = [l for l in lines if not l.startswith("## ")]
    # all the content under the first block, seven bare headings after it
    rec["Body"] = "\n".join([heads[0]] + rest + heads[1:])
    with pytest.raises(N.NarrativeRefusal, match=r"carries 0 word\(s\) against its floor"):
        N.write(wb, "assessment", "8", rec, actor="probe", card="REC-07", run=ready)


def test_a_block_under_the_docs_own_length_band_is_refused(ready):
    """§5's `What we see` is 350–550 words in the Doc; 100 words of it is
    not a deep dive, however long the other blocks run."""
    wb = ready.open()
    eids = _eids(ready)
    pillar = sorted({c[:2] for c in wb.selected_subcaps()})[0]
    rec = section_record("5", eids, "assessment")
    parts = N.blocks_split(rec["Body"])
    out = []
    for block, text in parts:
        if block == "What we see":
            text = " ".join(text.split()[:100])
        out.append(f"## {block}\n{text}\n")
    rec["Body"] = "\n".join(out)
    with pytest.raises(N.NarrativeRefusal, match="What we see.*floor of 350"):
        N.write(wb, "assessment", "5", rec, actor="probe", card=pillar, run=ready)


def test_a_recommendation_card_needs_its_own_title(ready):
    wb = ready.open()
    eids = _eids(ready)
    rec = section_record("8", eids, "assessment", Heading="Recommendations")
    with pytest.raises(N.NarrativeRefusal, match="own title"):
        N.write(wb, "assessment", "8", rec, actor="probe", card="REC-07", run=ready)
    rec = section_record("8", eids, "assessment", Heading="REC-07: Govern the member record")
    out = N.write(wb, "assessment", "8", rec, actor="probe", card="REC-07", run=ready)
    assert out["card"] == "REC-07"
    row = [r for r in wb.rows("Report_Narrative") if r.get("Card_ID") == "REC-07"][0]
    assert row["Heading"] == "Govern the member record"        # the id is the renderer's


def test_the_section_floor_counts_prose_not_table_rows(ready):
    wb = ready.open()
    eids = _eids(ready)
    rec = section_record("2", eids, "client_research")
    # a body that is all table: 600 rows of pipe cells under the blocks
    table = "\n".join(f"| P1C1.1.{i} | {2.0 + i / 100:.2f} | 3.0 | [{eids[0]}] |" for i in range(600))
    rec["Body"] = "\n".join(f"## {b}\n| Cell | Score | Median | Ev |\n|---|---|---|---|\n{table}\n"
                            for b in RS.SPECS["client_research"].section("2").blocks)
    with pytest.raises(N.NarrativeRefusal, match="prose words"):
        N.write(wb, "client_research", "2", rec, actor="probe", run=ready)


def test_a_paragraph_pasted_into_a_second_section_is_refused(ready):
    wb = ready.open()
    eids = _eids(ready)
    first = section_record("3", eids, "client_research")
    N.write(wb, "client_research", "3", first, actor="probe", run=ready)
    para = N.paragraphs_in(first["Body"])[0]
    second = section_record("7", eids, "client_research")
    second["Body"] = second["Body"].replace("## 7.2 Negative search results\n",
                                            "## 7.2 Negative search results\n" + para + "\n\n", 1)
    with pytest.raises(N.NarrativeRefusal, match="appear more than once"):
        N.write(wb, "client_research", "7", second, actor="probe", run=ready)


# ── the renderer keeps the Doc's shape ─────────────────────────────────────

def test_pipe_tables_in_a_body_render_as_word_tables_and_sheets_dump_once(tmp_path):
    run, wb, cells, ev = scored_run(tmp_path)
    eids = bank_evidence(wb, cells[0], n=7)
    for key in RS.SPECS:
        write_report(wb, key, eids, run=run)
    sign_off_sections(wb)
    spec = RS.SPECS["assessment"]
    out = R.render(wb, spec, tmp_path / "out")
    doc = Document(out["path"])
    # no paragraph is a pipe row
    assert not any(p.text.strip().startswith("|") for p in doc.paragraphs)
    heads = [c.text for t in doc.tables for c in t.rows[0].cells]
    assert "Steelman" in " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Data dependency" in " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    # each sheet extract once; the per-cell score sheet never
    titles = [p.text for p in doc.paragraphs if p.style.name == "Heading 2"]
    assert titles.count("Peer benchmarks") == 1, titles
    assert "Subcapability scores" not in titles
    assert "Evidence cited in this section" not in titles
    # a REC card is headed by its id and title, not the section heading
    assert any(re.match(r"REC-01: \S+ \S+", t) for t in titles), titles
    assert "Recommendations" not in titles
    # and the render says what was prose and what was table
    assert out["words"] > 0 and out["table_words"] > 0
    # the gate reads the same artefact and passes it
    assert GS.report_findings(out["path"], kind="assessment",
                              subcaps=len(cells), scores=GS._workbook_scores(wb.path)) == []
    # the Recommendations tab carries the card's title and no markup
    got = grains.recommendations(wb)
    rows = [r for r in wb.rows("Recommendations") if r.get("Rec_ID")]
    assert rows and all(r["Title"] != "Recommendations" for r in rows)
    assert all("## " not in str(r["Rationale"]) and "|---" not in str(r["Rationale"])
               for r in rows)


# ── the brief hands commands the CLI accepts ───────────────────────────────

def test_the_report_brief_names_commands_the_cli_accepts(ready):
    wb = ready.open()
    b = brief.report_batch(wb, run=ready, out_dir=ready.root / "07_qa" / "briefs_probe")
    import json
    for lane in b["briefs"]:
        pk = json.loads(Path(lane["prompt_file"]).with_suffix(".json").read_text())
        write = [c for c in pk["first_commands"] if "narrative write" in c][0]
        assert "--json" in write and "--body-file" not in write
        # every block carries the floor it owes
        assert all(re.search(r"\(\d+w\+", blk) for s in pk["sections"]
                   for blk in s.get("blocks", []))
        assert any("pipe rows" in r for r in pk["rules"])
    v = brief.report_batch(wb, run=ready, out_dir=ready.root / "07_qa" / "briefs_probe_v",
                           validator=True)
    pk = json.loads(Path(v["briefs"][0]["prompt_file"]).with_suffix(".json").read_text())
    review = [c for c in pk["first_commands"] if "narrative review" in c][0]
    m = re.search(r"--verdict (\S+)", review)
    assert set(m.group(1).split("|")) == set(N.VERDICTS), review


# ── the gold gate reads the docx, not its word count ───────────────────────

def _docx(path, items):
    import docx
    import zipfile
    d = docx.Document()
    for style, text in items:
        if style == "TABLE":
            tb = d.add_table(rows=0, cols=len(text[0]))
            for row in text:
                cells = tb.add_row().cells
                for i, v in enumerate(row):
                    cells[i].text = str(v)
        else:
            d.add_paragraph(text, style=style)
    d.save(str(path))
    with zipfile.ZipFile(str(path), "a") as z:
        z.writestr("word/header1.xml", "<hdr/>")
    return path


def test_the_gold_gate_catches_pasted_paragraphs_and_cards_without_tables(tmp_path):
    items = [("Heading 1", h) for h in RS.numbered_headings("assessment")[:1]]
    items += [("Normal", OLD_FILLER)] * 3
    items += [("Heading 1", "8. Recommendations"), ("Heading 2", "REC-01: do a thing"),
              ("Heading 3", "Rebuttal"), ("Normal", "It survives.")]
    items += [("Heading 2", "5.1 Pillar deep dive (P1): a pillar"),
              ("Normal", "AI and data overlay " + "word " * 60)]
    codes = [f["code"] for f in GS.report_findings(
        _docx(tmp_path / "DMA_Assessment_Report_x.docx", items), kind="assessment")]
    assert "GS-RPT-BOILERPLATE" in codes
    assert "GS-RPT-TABLES" in codes
    assert "GS-RPT-REBUTTALS" in codes


def test_the_gold_length_floor_counts_prose_not_tables(tmp_path):
    items = [("Heading 1", h) for h in RS.numbered_headings("assessment")]
    items.append(("TABLE", [["a"] * 4] + [["word"] * 4 for _ in range(3000)]))   # 12,000 table words
    f = GS.report_findings(_docx(tmp_path / "DMA_Assessment_Report_x.docx", items),
                           kind="assessment")
    assert any(x["code"] == "GS-RPT-LENGTH" and "prose words" in x["detail"] for x in f)


def test_the_package_gate_reconciles_the_workbooks_scores_into_the_report(tmp_path):
    items = [("Heading 1", h) for h in RS.numbered_headings("assessment")]
    items.insert(1, ("Normal", "The overall score is 2.12 (M2); P1 sits at 2.12."))
    path = _docx(tmp_path / "DMA_Assessment_Report_x.docx", items)
    ok = GS.report_findings(path, kind="assessment",
                            scores={"overall": 2.12, "pillars": {"P1": 2.12}})
    assert not [x for x in ok if x["code"] == "GS-RPT-RECONCILE"]
    bad = GS.report_findings(path, kind="assessment",
                             scores={"overall": 2.40, "pillars": {"P1": 2.12}})
    assert [x for x in bad if x["code"] == "GS-RPT-RECONCILE"]


def test_the_package_gate_reads_scores_from_the_workbook(ready):
    got = GS._workbook_scores(ready.workbook_path)
    assert got and got["overall"] is not None and got["pillars"]
