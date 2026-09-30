"""A report can pass every volume floor and still be the wrong SHAPE.

THE DEFECT THESE TESTS PIN, measured 2026-09-06 on a delivered pair of reports
that the gate of the day passed with zero findings (GSY-31):

    assessment   50 tables against the reference's 92, paragraph words 1.41x
    research     26 tables against the reference's 39, paragraph words 1.33x

The owner's report was "the 2 reports lack depth ... do not adhere to template
requirements eg where tables are, I see paragraphs." The gate counted words,
citations and section headings, none of which can tell a report that TABULATES
its register from one that describes it in a paragraph — and because more prose
raises a word count, the defect was not merely uncaught, it was REWARDED.

Three things are asserted here: the whole-report structure gates catch the
shape (GS-RPT-TABLE-FLOOR, GS-RPT-PROSE-FOR-STRUCTURE, GS-RPT-TABLE-DUMP), the
renderer predicts that shape BEFORE a paragraph is written by the same grammar
it renders with (`reports._predicted_shape` walks `_emit_lines`' own rules, so
the prediction and the post-render measurement agree), and the anti-patterns
reach the writer before it writes (engine/authoring.py) rather than only the
gate after it ships.
"""
from __future__ import annotations

import json
import zipfile

import pytest

from engine import authoring as A
from engine import gold_standard as GS
from engine import report_spec as RS
from engine import reports as R
from engine import template as T

from fixtures import bank_evidence, scored_run, sign_off_sections, write_report


@pytest.fixture(scope="module")
def gold() -> dict:
    return json.loads((T.TEMPLATES_DIR / "gold_reference.json").read_text())


# ── the calibration discipline, extended to the structure floors ──────────

def test_the_table_floor_is_at_or_below_what_golden_1_meets(gold):
    """The rule gold_standard.py has stated since it was written: a floor the
    reference itself would fail is a floor nobody measured."""
    for kind in ("research", "assessment"):
        floors = GS.depth_floors(kind)
        assert floors["tables"] <= gold["reports"][kind]["tables"], (kind, floors)


def test_the_prose_and_dump_limits_do_not_refuse_the_reference(gold):
    for kind in ("research", "assessment"):
        ref = gold["reports"][kind]
        floors = GS.depth_floors(kind)
        # The reference's own prose sits inside the inflation limit …
        assert ref["words_paragraphs"] <= (
            floors["reference_paragraph_words"] * GS.PROSE_INFLATION_LIMIT)
        # … and its own average table is, by construction, 1x its average.
        assert GS._gold_avg_table_words(kind) >= 1


def test_the_table_floor_scales_with_the_run_and_never_vanishes():
    small = GS.depth_floors("assessment", subcaps=6)
    full = GS.depth_floors("assessment")
    assert 1 <= small["tables"] < full["tables"]


# ── fixtures: a report of a chosen shape ──────────────────────────────────

def _report(path, *, tables, rows_per_table, prose_words, kind="assessment"):
    """A synthetic report carrying an exact number of tables of an exact size,
    and an exact quantity of prose — so a shape finding can be provoked or
    avoided on purpose."""
    import docx
    d = docx.Document()
    for h in RS.numbered_headings(kind if kind == "assessment" else "client_research"):
        d.add_paragraph(h, style="Heading 1")
    # citations, so the volume gates are not what fires
    d.add_paragraph(" ".join(f"[E-{i}]" for i in range(1, 200)))
    d.add_paragraph(" ".join(["word"] * prose_words))
    for _ in range(tables):
        t = d.add_table(rows=1, cols=3)
        for i, c in enumerate(("A", "B", "C")):
            t.rows[0].cells[i].text = c
        for _ in range(rows_per_table):
            cells = t.add_row().cells
            for i in range(3):
                cells[i].text = "cell value here"
    d.save(str(path))
    with zipfile.ZipFile(str(path), "a") as z:
        z.writestr("word/header1.xml", "<hdr/>")
    return path


def _codes(path, kind="assessment", subcaps=690):
    return {f["code"] for f in GS.report_findings(path, kind=kind, subcaps=subcaps)}


# ── detection ─────────────────────────────────────────────────────────────

def test_prose_where_a_table_belongs_is_refused(tmp_path):
    """Few tables, inflated prose — the delivered pair's exact shape."""
    p = _report(tmp_path / "thin.docx", tables=50, rows_per_table=2,
                prose_words=20000)
    codes = _codes(p)
    assert "GS-RPT-TABLE-FLOOR" in codes
    assert "GS-RPT-PROSE-FOR-STRUCTURE" in codes


def test_a_sheet_emitted_whole_is_refused(tmp_path):
    """Enough tables, but each one is a pasted sheet rather than an extract."""
    p = _report(tmp_path / "dump.docx", tables=92, rows_per_table=400,
                prose_words=11000)
    assert "GS-RPT-TABLE-DUMP" in _codes(p)


def test_a_report_of_the_reference_shape_passes_every_structure_gate(tmp_path):
    """The floors must be reachable: a report built to Golden 1's own measured
    shape raises none of the three structure findings."""
    p = _report(tmp_path / "good.docx", tables=92, rows_per_table=3,
                prose_words=11000)
    codes = _codes(p)
    for code in ("GS-RPT-TABLE-FLOOR", "GS-RPT-PROSE-FOR-STRUCTURE",
                 "GS-RPT-TABLE-DUMP"):
        assert code not in codes, (code, codes)


def test_long_prose_alone_is_not_refused_when_the_tables_are_there(tmp_path):
    """The finding is prose STANDING IN FOR structure, not prose. A section
    that tabulates what it should and still argues at length is not the
    defect, and a gate that punished it would push producers to cut analysis."""
    p = _report(tmp_path / "rich.docx", tables=92, rows_per_table=3,
                prose_words=30000)
    assert "GS-RPT-PROSE-FOR-STRUCTURE" not in _codes(p)


def test_the_whole_report_floor_and_the_per_card_floor_are_different_gates(tmp_path):
    """GS-RPT-TABLES is the per-card floor (three under a REC, two under a
    deep dive); GS-RPT-TABLE-FLOOR is the whole report's count against the
    reference density. A report with no cards at all owes the second and
    not the first."""
    p = _report(tmp_path / "nocards.docx", tables=3, rows_per_table=2,
                prose_words=11000)
    codes = _codes(p)
    assert "GS-RPT-TABLE-FLOOR" in codes
    assert "GS-RPT-TABLES" not in codes


# ── the prediction walks the renderer's own grammar ───────────────────────

def test_the_predicted_shape_reads_a_body_the_way_emit_lines_renders_it():
    body = ("## Capability scorecard\n\n"
            "| Cell | Score | Peer |\n"
            "|---|---|---|\n"
            "| P4C1.2.1 | 1.6 | 3.0 |\n"
            "| P4C1.2.2 | 1.25 | 3.0 |\n\n"
            "The scorecard above states the two cells this rests on.\n"
            "| Second | table |\n"
            "| directly | after prose |\n"
            "and prose directly after it.\n\n"
            "|---|---|\n"
            "The estate splits OPS | CUST | DATA across four layers.")
    tables, prose, sizes = R._predicted_shape(body)
    # two tables: the scorecard, and the two-row table between the prose
    # lines; a separator-only run is no table; a pipe INSIDE a sentence is prose
    assert tables == 2
    assert sizes == [R._gs_words("Cell Score Peer P4C1.2.1 1.6 3.0 P4C1.2.2 1.25 3.0"),
                     R._gs_words("Second table directly after prose")]
    assert prose == R._gs_words(
        "Capability scorecard The scorecard above states the two cells this "
        "rests on. and prose directly after it. The estate splits OPS CUST "
        "DATA across four layers.")
    # and the same body rendered carries exactly those tables
    import docx
    d = docx.Document()
    R._emit_body(d, body)
    assert len(d.tables) == tables
    assert [sum(R._gs_words(c.text) for r in t.rows for c in r.cells)
            for t in d.tables] == sizes
    assert sum(R._gs_words(p.text) for p in d.paragraphs) == prose


def test_the_pre_render_prediction_matches_the_rendered_measurement(tmp_path):
    """`reports.check` refuses on the SAME numbers `gold_standard` measures
    from the rendered file — tables and table words exactly, and prose no
    higher than the render (the cover note, the scope lines and the sources
    list are render-time additions the prediction does not claim)."""
    run, wb, cells, ev = scored_run(tmp_path)
    eids = bank_evidence(wb, cells[0], n=7)
    for key in RS.SPECS:
        write_report(wb, key, eids, run=run)
    sign_off_sections(wb)
    for key in RS.SPECS:
        spec = RS.SPECS[key]
        predicted = R._predicted_report_shape(wb, R.curate(wb, spec))
        out = R.render(wb, spec, tmp_path / "out")
        shape = GS._docx_shape(out["path"])
        assert predicted["tables"] == shape["tables"], (key, predicted, shape)
        assert predicted["table_words"] == shape["table_words"], (key, predicted, shape)
        assert predicted["largest_table"] == shape["largest_table"], key
        assert predicted["paragraph_words"] <= shape["paragraph_words"], key
        assert predicted["tables"] >= GS.depth_floors(
            "assessment" if key == "assessment" else "research",
            subcaps=len(cells))["tables"]


def test_the_projected_recommendations_tab_does_not_reprint_the_cards(tmp_path):
    """The pipeline projects the Recommendations tab from §8's cards before
    it renders (`grains.recommendations`), and the tab's Rationale is each
    card's whole argument. Rendered whole, §9's extract re-emitted §8 as one
    ~1,800-word table — 40% of the report's table content, the dump shape —
    and the pre-render check refused the walk's own report (measured
    30-09-2026 on stress_stage_and_supersede). The extract is the register's
    columns; the argument stays in the card."""
    from docx import Document
    from engine import grains
    run, wb, cells, ev = scored_run(tmp_path)
    eids = bank_evidence(wb, cells[0], n=7)
    write_report(wb, "assessment", eids, run=run)
    sign_off_sections(wb)
    grains.recommendations(wb)
    rows = [r for r in wb.rows("Recommendations") if r.get("Rec_ID")]
    assert rows and all(len(str(r.get("Rationale") or "").split()) > 100 for r in rows)
    spec = RS.SPECS["assessment"]
    assert not [p for p in R.check(wb, R.curate(wb, spec)) if "GS-RPT-TABLE-DUMP" in p]
    out = R.render(wb, spec, tmp_path / "out")
    doc = Document(out["path"])
    heads = [[c.text for c in t.rows[0].cells] for t in doc.tables]
    rec = [h for h in heads if h and h[0] == "Rec_ID"]
    assert rec == [["Rec_ID", "Title", "Category_ID", "Priority", "Horizon", "Owner"]], heads
    codes = {f["code"] for f in GS.report_findings(out["path"], kind="assessment",
                                                    subcaps=len(cells))}
    assert "GS-RPT-TABLE-DUMP" not in codes
    assert GS._docx_shape(out["path"])["tables"] == \
        R._predicted_report_shape(wb, R.curate(wb, spec))["tables"]


def test_the_renderer_refuses_a_report_short_of_its_table_floor(tmp_path, monkeypatch):
    """The pre-render half of GS-RPT-TABLE-FLOOR: strip every authored table
    out of every body and every declared sheet's contribution, hold the run
    to a full-size floor, and `reports.check` names the floor before a
    paragraph is written — with the prose-for-structure reading beside it."""
    run, wb, cells, ev = scored_run(tmp_path)
    eids = bank_evidence(wb, cells[0], n=7)
    write_report(wb, "assessment", eids, run=run)
    spec = RS.SPECS["assessment"]
    curated = R.curate(wb, spec)
    assert not any("GS-RPT-TABLE-FLOOR" in p for p in R.check(wb, curated))
    for b in curated["blocks"]:
        b["body"] = "\n".join(l for l in b["body"].splitlines()
                              if not l.strip().startswith("|"))
        b["tables"] = []
    full = GS.depth_floors("assessment")          # the reference's own 690 cells
    monkeypatch.setattr(GS, "depth_floors", lambda kind, subcaps=None: full)
    problems = R.check(wb, curated)
    assert any("GS-RPT-TABLE-FLOOR" in p for p in problems), problems
    assert any("GS-RPT-PROSE-FOR-STRUCTURE" in p for p in problems), problems


# ── the anti-patterns reach the writer BEFORE it writes ───────────────────

def test_every_antipattern_names_a_gate_and_a_repair():
    pats = A.antipatterns()
    assert pats, "the anti-pattern register must not be empty"
    for a in pats:
        for field in ("id", "gate", "what", "measured", "repair"):
            assert a.get(field), (a.get("id"), field)


def test_the_register_covers_every_structure_gate():
    """A gate with no entry in the register is a rule the writer is never told
    about until it refuses them."""
    gates = {a["gate"] for a in A.antipatterns()}
    for code in ("GS-RPT-TABLE-FLOOR", "GS-RPT-PROSE-FOR-STRUCTURE",
                 "GS-RPT-TABLE-DUMP", "GS-RPT-COVER", "GS-RPT-FRONTMATTER",
                 "GS-RPT-SECTION-DISTRIBUTION", "GS-RPT-DEGENERATE-TABLE",
                 "GS-RPT-PROSE-DUMP"):
        assert code in gates, (code, gates)


def test_every_gate_the_register_names_exists_in_the_engine():
    """The register may not name a gate code the gate module does not emit
    (a connector gate, CG-NN, lives in the app and is named as such)."""
    import inspect
    src = inspect.getsource(GS)
    for a in A.antipatterns():
        if a["gate"].startswith("GS-"):
            assert f'"{a["gate"]}"' in src, a["gate"]


def test_the_brief_binds_the_antipatterns_to_this_sections_own_tables():
    """A general rule is advice; the section's own declared inputs are an
    instruction. §6 of the assessment declares the technology estate sheets."""
    text = A.brief("assessment", "6")
    assert "Benchmark and Technology Estate" in text
    assert "TABLES THIS SECTION OWES" in text
    assert "GOLDEN 1 CARRIES 7 TABLE(S) IN THIS SECTION" in text
    for sheet in A.table_inputs(
            next(s for s in RS.SPECS["assessment"].sections if str(s.id) == "6")):
        assert sheet in text
    for a in A.antipatterns():
        assert a["id"] in text


def test_the_briefs_table_obligations_are_the_renderers_own():
    """`table_inputs` applies the renderer's own `_NO_TABLE` filter, so the
    sheet a writer is told to state is a sheet the renderer will look for —
    and the per-cell score sheet and the register are never owed."""
    for spec in RS.SPECS.values():
        for sec in spec.sections:
            owed = A.table_inputs(sec)
            assert not set(owed) & R._NO_TABLE, (spec.key, sec.id, owed)
            assert set(owed) <= set(sec.inputs)


class _WB:
    def __init__(self, empty=()):
        self.empty = set(empty)

    def rows(self, name):
        return [] if name in self.empty else [{"x": 1}]

    def selected_subcaps(self):
        return ["P1C1.1.1"] * 6


def test_the_brief_names_an_empty_declared_input_as_the_thing_to_fix_first():
    """An input that carries no rows renders no table, and the section then
    'lacks a table' for a reason no amount of writing can repair."""
    text = A.brief("assessment", "6", _WB(empty={"Tech_Register"}))
    assert "Tech_Register" in text
    assert "EMPTY" in text
    ob = A.section_obligations("assessment", "6", _WB(empty={"Tech_Register"}))
    assert "Tech_Register" in ob["tables_empty"]
    assert "Tech_Register" not in ob["tables_that_will_render"]


def test_preflight_reports_the_whole_reports_table_position():
    out = A.preflight("assessment", _WB())
    assert out["tables_declared"] > 0
    assert out["tables_that_will_render"] == out["tables_declared"]
    assert out["empty_declared_inputs"] == []
    assert len(out["sections"]) == len(RS.SPECS["assessment"].sections)


def test_the_brief_step_reaches_the_producer_agents():
    """The generated producers carry the brief as a first step, so the rule
    reaches the writer through the manifest and not only through a gate."""
    from pathlib import Path
    agents = (Path(__file__).resolve().parents[3] / "plugins" / "dma-insights"
              / "agents" / "reports")
    for name in ("report-research-producer.md", "report-assessment-producer.md"):
        text = (agents / name).read_text(encoding="utf-8")
        assert "engine.authoring brief" in text, name
        assert "report_antipatterns.json" in text, name
