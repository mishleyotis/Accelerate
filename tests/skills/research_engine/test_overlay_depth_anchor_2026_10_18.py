"""An overlay block starts at its heading, not at every mention of the phrase
(EPIC, 2026-10-18: "P4 is where the AI and data overlay concentrates" in the
block's own prose cut a ~247-word overlay to 8 words and failed the gate)."""
from docx import Document

from engine import gold_standard as GS

FILLER = ("The firm runs a claims model on its lakehouse and staffs a data team, "
          "and the readiness of each workload is cited to the register [E-001]. ")


def _report(tmp_path, prose_mentions: bool):
    d = Document()
    for p in ("P1", "P2", "P3", "P4"):
        d.add_heading(f"5.{p[1]} Pillar deep dive ({p}): pillar", level=2)
        d.add_heading("AI and data overlay", level=3)
        lead = (f"{p} is where the AI and data overlay concentrates: " if prose_mentions
                else f"{p} carries the overlay: ")
        d.add_paragraph(lead + FILLER * 6)
    path = tmp_path / "DMA_Assessment_Report_test.docx"
    d.save(str(path))
    return path


def _depth(findings):
    return [f for f in findings if f["code"] == "GS-RPT-AIOVERLAY-DEPTH"]


def test_a_mention_inside_the_block_does_not_split_it(tmp_path):
    assert not _depth(GS.report_findings(_report(tmp_path, True), kind="assessment"))


def test_a_thin_block_under_its_heading_still_fails(tmp_path):
    d = Document()
    for p in ("P1", "P2", "P3", "P4"):
        d.add_heading(f"5.{p[1]} Pillar deep dive ({p}): pillar", level=2)
        d.add_heading("AI and data overlay", level=3)
        d.add_paragraph("Models.")
    path = tmp_path / "DMA_Assessment_Report_thin.docx"
    d.save(str(path))
    assert _depth(GS.report_findings(path, kind="assessment"))
