"""Executive_Summary states the client-facing overall, and the worker reads it
(MEM-0561; SWBC gold audit 2026-10-04, stream P5).

Measured on SWBC (run 7968492e): the scoring workbook states the overall
maturity in its `Executive_Summary` tab — "Overall Maturity", pillar-weighted
20/20/30/30 = 2.0073. The producers never opened the tab, and the map a
producer consults to decide which tabs feed surfaces,
`workbook_tab_coverage`, classified it ('run config', 'not_client_facing')
— telling them the one tab stating the client-facing headline number was not
client-facing. The parser's own docstring already recorded that
Executive_Summary is one of four places the overall is stated and that "no
reader claimed any of them" on Golden 1.

These pin both halves of the fix:
- the tab map names Executive_Summary as a source for overview.scores and
  overview.exec_summary, and a reader claims it;
- the parser READS its "Overall Maturity" as one of the composite's stated
  sources — at the precision the workbook states, never rounded — fills the
  composite from it when the grain tabs state none, and records a
  disagreement observation when two stated sources differ.
"""
from decimal import Decimal

import openpyxl
import pytest


def _exec_summary(wb, value=2.0073, label="Overall Maturity", first=False):
    ws = wb.active if first else wb.create_sheet("Executive_Summary")
    ws.title = "Executive_Summary"
    ws.append(["SWBC — Digital Maturity Assessment"])
    ws.append([])
    ws.append(["Metric", "Value", "Basis"])
    ws.append([label, value, "pillar-weighted 20/20/30/30"])
    ws.append(["Band", "Building", None])
    return ws


def _pillar_summary(wb, overall=None):
    ws = wb.create_sheet("Pillar_Summary")
    ws.append(["Pillar", "Name", "Weighted_Score"])
    ws.append(["P1", "Strategy", 1.85])
    ws.append(["P2", "Experience", 2.10])
    ws.append(["P3", "Operations", 2.04])
    ws.append(["P4", "Data", 2.02])
    if overall is not None:
        ws.append(["OVERALL", "weighted", overall])
    return ws


def _subcap(wb):
    ws = wb.create_sheet("P1_Subcap_Scoring")
    ws.append(["SubCap_ID", "SubCap_Name", "Effective_Score"])
    ws.append(["P1C1.1.1", "A capability", 2.0])


def _save(wb, tmp_path, name="wb.xlsx"):
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]
    path = tmp_path / name
    wb.save(path)
    return str(path)


def test_the_tab_map_calls_executive_summary_client_facing():
    from dma_worker.workbook_parser import _TAB_TARGET, _TAB_READERS
    feeds, confidence = _TAB_TARGET["Executive_Summary"]
    assert confidence != "not_client_facing", (
        "Executive_Summary states the headline composite — MEM-0561")
    assert "overview.scores" in feeds and "overview.exec_summary" in feeds
    assert "Executive_Summary" in _TAB_READERS, "a reader must claim it"


def test_the_census_reports_it_read_not_run_config(tmp_path):
    from dma_worker.workbook_parser import workbook_tab_coverage
    wb = openpyxl.Workbook()
    _exec_summary(wb, first=True)
    cov = workbook_tab_coverage(_save(wb, tmp_path))
    assert "Executive_Summary" not in cov["unread_with_rows"]
    assert cov["tabs_read"] == 1


def test_the_stated_overall_is_read_when_the_grain_tab_states_none(tmp_path):
    """SWBC's shape: the grain tab carries four pillars and no OVERALL row,
    and Executive_Summary states 2.0073. The composite is READ from it, at
    the stated precision (invariant 6 bands the RAW score)."""
    from dma_worker.workbook_parser import parse_scoring_workbook
    wb = openpyxl.Workbook()
    _pillar_summary(wb, overall=None)
    _subcap(wb)
    _exec_summary(wb)
    out = parse_scoring_workbook(_save(wb, tmp_path))
    assert out.composite == Decimal("2.0073")
    assert out.composite_source_cell == "Executive_Summary!B4"


def test_the_grain_tab_still_wins_and_a_disagreement_is_recorded(tmp_path):
    from dma_worker.workbook_parser import parse_scoring_workbook
    wb = openpyxl.Workbook()
    _pillar_summary(wb, overall=2.25)
    _subcap(wb)
    _exec_summary(wb, value=2.0073)
    out = parse_scoring_workbook(_save(wb, tmp_path))
    assert float(out.composite) == pytest.approx(2.25)
    assert out.composite_source_cell == "Pillar_Summary!C6"
    hits = [o for o in out.observations if o.kind == "stated_overall_disagreement"]
    assert len(hits) == 1
    d = hits[0].detail
    assert {r["source_cell"] for r in d["readings"]} == {
        "Pillar_Summary!C6", "Executive_Summary!B4"}


def test_agreeing_stated_sources_record_nothing(tmp_path):
    from dma_worker.workbook_parser import parse_scoring_workbook
    wb = openpyxl.Workbook()
    _pillar_summary(wb, overall=2.0073)
    _subcap(wb)
    _exec_summary(wb, value=2.0073)
    out = parse_scoring_workbook(_save(wb, tmp_path))
    assert not [o for o in out.observations if o.kind == "stated_overall_disagreement"]


def test_an_executive_summary_without_an_overall_is_not_a_composite(tmp_path):
    """Absent beats invented: a band word or a prose cell is not a figure."""
    from dma_worker.workbook_parser import parse_scoring_workbook
    wb = openpyxl.Workbook()
    _pillar_summary(wb, overall=None)
    _subcap(wb)
    _exec_summary(wb, value="Building", label="Overall Maturity")
    out = parse_scoring_workbook(_save(wb, tmp_path))
    assert out.composite is None and out.composite_source_cell is None
