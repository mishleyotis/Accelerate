"""The first-call deck renders the app's bands and the assessment's levels.

Measured 28-09-2026 (QA audit F-L14-041): the deck cut bands at
1.50 / 2.50 / 3.50 (the app cuts strictly at 2 / 3 / 4 on the raw score),
painted the retired fifth-band hex on Differentiating, labelled the fifth
maturity level with the banned band word, and required a score for a
category the v7.0 catalogue does not have. One owner each now:
apps/web/lib/bands.js for the band fills, engine.contract.band_of for the
boundaries, engine/rubric.py for the 1-5 score levels. This file pins the
deck's colour module — and the heatmap editor that writes from it — to all
three.
"""
import importlib.util
import re
import sys
from pathlib import Path

import pytest

REPO = Path(__file__).resolve().parents[4]
PLUGIN = REPO / "plugins" / "dma-insights"
DECK = PLUGIN / "skills" / "dma-first-call-deck"
BANDS_JS = REPO / "apps" / "web" / "lib" / "bands.js"

sys.path.insert(0, str(PLUGIN / "skills" / "dma-research"))
from engine import contract, rubric  # noqa: E402


def _load(path, name):
    d = str(path.parent)
    if d not in sys.path:
        sys.path.insert(0, d)
    spec = importlib.util.spec_from_file_location(name, path)
    m = importlib.util.module_from_spec(spec)
    sys.modules[name] = m
    spec.loader.exec_module(m)
    return m


cls = _load(DECK / "references" / "01_brand" / "color_level_system.py",
            "color_level_system")


def _app_fills():
    return dict(re.findall(r'(\w+): \{ fill: "#([0-9A-F]{6})"', BANDS_JS.read_text()))


# ── bands: fills from bands.js, boundaries from contract.band_of ──────────

def test_the_deck_accents_are_the_apps_band_fills():
    fills = _app_fills()
    assert set(fills) == set(contract.BANDS)
    for band in contract.BANDS:
        assert cls.LEVEL_4TIER[band]["accent"] == fills[band], band


@pytest.mark.parametrize("score", [0.0, 0.5, 1.49, 1.5, 1.99, 2.0, 2.49, 2.5,
                                   2.99, 3.0, 3.49, 3.5, 3.99, 4.0, 4.49, 4.5, 5.0])
def test_the_deck_bands_where_the_app_bands(score):
    assert cls.score_to_level_4tier(score) == contract.band_of(score)


def test_the_deck_ranges_start_where_the_bands_start():
    starts = [cls.LEVEL_4TIER[b]["score_range"][0] for b in contract.BANDS]
    assert starts == [0.0, *contract._BAND_MAX]


def test_the_legends_show_the_fills_the_bars_use():
    a = {b: cls.LEVEL_4TIER[b]["accent"] for b in contract.BANDS}
    s = cls.STATIC_COLORS
    assert (s["s14_legend_act_acc"], s["s14_legend_bld_acc"],
            s["s14_legend_cmp_acc"], s["s14_legend_dif_acc"]) == tuple(a.values())
    assert (s["s10_legend_act_acc"], s["s10_legend_bld_acc"],
            s["s10_legend_cmp_acc"]) == tuple(a.values())[:3]
    # the Slide 10 legend swatches are WRITTEN, so a template's stale swatch
    # is repainted rather than merely reported
    for idx in (27, 30, 33):
        role = cls.ALL_SLIDE_ROLES[10][idx]
        assert role[1] == "static" and role[4] is True and role[5] is True, idx


# ── levels: names and cuts from engine/rubric.py ─────────────────────────

def test_the_slide13_levels_are_the_rubrics():
    assert [cls.LEVEL_5TIER[n]["label"] for n in sorted(cls.LEVEL_5TIER)] == \
        [r[1] for r in rubric.RUBRIC]
    for score in [0.0, 1.0, 1.49, 1.5, 2.49, 2.5, 3.49, 3.5, 4.49, 4.5, 5.0]:
        assert f"M{cls.score_to_level_5tier(score)}" == rubric.maturity_level(score), score


# ── a retired capability renders unscored, never invented ────────────────

def test_a_null_score_has_no_band_and_no_maturity_fill():
    assert cls.score_to_level_4tier(None) is None
    assert contract.band_of(None) is None
    role = ("accent_strip", "data", "s14.scores[0]", "accent", True, True, False, False)
    hexv = cls.get_expected_hex(role, {"s14": {"scores": [None]}}, slide_num=14)
    assert hexv == cls.UNSCORED["accent"]
    assert hexv not in {d["accent"] for d in cls.LEVEL_4TIER.values()}
    assert "Sustainable Finance & ESG" in cls.RETIRED_CAPABILITIES
    assert set(cls.RETIRED_CAPABILITIES) <= set(cls.CAPABILITY_ORDER)


def _heatmap_editor():
    return _load(DECK / "scripts" / "03_editing" / "heatmap_editor.py", "heatmap_editor")


def test_the_editor_accepts_null_only_for_a_retired_capability():
    he = _heatmap_editor()
    full = {c: 2.5 for c in cls.CAPABILITY_ORDER}
    assert he.validate_inputs(full, full) == []
    no_esg = {c: 2.5 for c in cls.CAPABILITY_ORDER if c not in cls.RETIRED_CAPABILITIES}
    assert he.validate_inputs(no_esg, no_esg) == []
    nulled = dict(full, **{"Sustainable Finance & ESG": None})
    assert he.validate_inputs(nulled, nulled) == []
    wrong = dict(full, **{"Data Governance": None})
    errs = he.validate_inputs(wrong, full)
    assert errs and "not retired" in errs[0]
    assert he.compute_bar_width_emu(None) == 0
    assert he.compute_median_x_emu(1000, None) == 1000


def _block_slide():
    """One 8-shape heatmap block at base 0, shaped like the template's."""
    from pptx import Presentation
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    sh = slide.shapes
    sh.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(3), Inches(1))      # +0 card
    sh.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.1), Inches(1))    # +1 strip
    sh.add_textbox(Inches(0.2), Inches(0), Inches(2), Inches(0.3)).text_frame.text = "Cap"  # +2
    sh.add_textbox(Inches(2.5), Inches(0), Inches(0.5), Inches(0.3)).text_frame.text = "3.1"  # +3
    sh.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(0.5), Inches(2), Inches(0.1))  # +4 track
    sh.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.2), Inches(0.5), Inches(1), Inches(0.1))  # +5 bar
    sh.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(1), Inches(0.45), Inches(1), Inches(0.65))  # +6
    sh.add_textbox(Inches(0.2), Inches(0.7), Inches(2), Inches(0.3)).text_frame.text = "COMPETING"  # +7
    return list(slide.shapes)


def _roles_for_block_zero():
    out = {}
    for off, (name, ctype, source, key, wf, wb, wt, wtc) in cls.SLIDE_14_BLOCK_OFFSETS.items():
        src = "s14.scores[0]" if ctype == "data" else source
        out[off] = (f"{name}_cap01", ctype, src, key, wf, wb, wt, wtc)
    return out


def _edit(score, median):
    he = _heatmap_editor()
    ec = sys.modules["_editor_common"]
    shapes = _block_slide()
    audit = ec.EditorAudit(slide_num=14, editor_name="test")
    data = {"s14": {"scores": [score], "medians": [median]}, "input": {}}
    he.edit_block(shapes, 0, 0, "Digital Strategy & Vision", score, median,
                  data, audit, _roles_for_block_zero())
    return shapes, ec, audit


def test_the_editor_writes_the_apps_fill_for_a_scored_block():
    shapes, ec, audit = _edit(1.2, 2.0)
    assert ec.read_shape_fill_hex(shapes[1]).upper() == _app_fills()["Activating"]
    assert ec.read_shape_fill_hex(shapes[5]).upper() == _app_fills()["Activating"]
    assert shapes[7].text_frame.text == "ACTIVATING"
    assert shapes[3].text_frame.text == "1.2"
    assert shapes[5].width > 0
    assert not audit.errors


def test_the_editor_renders_a_retired_block_as_not_assessed():
    shapes, ec, audit = _edit(None, None)
    assert shapes[7].text_frame.text == cls.UNSCORED["label"]
    assert shapes[3].text_frame.text == cls.UNSCORED["score_text"]
    assert ec.read_shape_fill_hex(shapes[5]).upper() == cls.UNSCORED["accent"]
    assert ec.read_shape_fill_hex(shapes[0]).upper() == cls.UNSCORED["card_bg"]
    assert shapes[5].width == 0
    fills = {d["accent"] for d in cls.LEVEL_4TIER.values()}
    assert ec.read_shape_fill_hex(shapes[1]).upper() not in fills
    assert not audit.errors


# ── the retired hex and the fifth band word are gone from the deck skill ──

def test_no_retired_hex_or_fifth_band_word_in_the_deck_skill():
    hits = []
    for p in DECK.rglob("*"):
        if p.suffix not in (".py", ".md", ".yaml", ".json") or \
                "deprecated" in p.parts or "__pycache__" in p.parts:
            continue
        for n, line in enumerate(p.read_text(errors="ignore").splitlines(), 1):
            if "185F60" in line.upper() or re.search(r"\bTransformational\b", line):
                hits.append(f"{p.relative_to(DECK)}:{n}")
    assert not hits, hits
